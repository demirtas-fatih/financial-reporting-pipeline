"""
pipeline.py
Automated Financial Reporting Pipeline
--------------------------------------
Reads raw financial CSV data, cleans it, computes KPIs,
and produces a formatted Excel report with charts.

Author : Fatih Demirtas
GitHub : https://github.com/demirtas-fatih
"""

import pandas as pd
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.ticker as mticker
import openpyxl
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, numbers
)
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from datetime import datetime
from pathlib import Path
import warnings
warnings.filterwarnings("ignore")

# ── Config ────────────────────────────────────────────────────────────────────
INPUT_FILE  = Path("data/financial_data.csv")
OUTPUT_DIR  = Path("output")
OUTPUT_DIR.mkdir(exist_ok=True)
REPORT_FILE = OUTPUT_DIR / f"Financial_Report_{datetime.today().strftime('%Y-%m-%d')}.xlsx"

# Brand colours
BLUE_DARK   = "1E3A5F"
BLUE_MID    = "2563EB"
BLUE_LIGHT  = "DBEAFE"
ACCENT      = "06B6D4"
GREEN       = "16A34A"
RED_SOFT    = "DC2626"
GREY_LIGHT  = "F8FAFC"
GREY_MID    = "E2E8F0"
WHITE       = "FFFFFF"

# ── 1. LOAD & CLEAN ───────────────────────────────────────────────────────────
print("📂 Loading data …")
df = pd.read_csv(INPUT_FILE, parse_dates=["date"])

# Drop unapproved rows & nulls
df = df[df["approved"] == True].dropna()

# Derived columns
df["month"]      = df["date"].dt.to_period("M")
df["variance"]   = df["actual"] - df["budget"]
df["var_pct"]    = ((df["actual"] - df["budget"]) / df["budget"] * 100).round(2)
df["over_budget"]= df["actual"] > df["budget"]

print(f"   ✅ {len(df):,} approved transactions loaded.")

# ── 2. AGGREGATE ──────────────────────────────────────────────────────────────
print("🔢 Computing KPIs …")

# Monthly summary
monthly = (
    df.groupby("month")
      .agg(budget=("budget","sum"), actual=("actual","sum"), txn=("actual","count"))
      .reset_index()
)
monthly["month_str"] = monthly["month"].astype(str)
monthly["variance"]  = monthly["actual"] - monthly["budget"]
monthly["var_pct"]   = ((monthly["variance"] / monthly["budget"]) * 100).round(2)

# Department summary
by_dept = (
    df.groupby("department")
      .agg(budget=("budget","sum"), actual=("actual","sum"), txn=("actual","count"))
      .reset_index()
      .sort_values("actual", ascending=False)
)
by_dept["var_pct"] = ((by_dept["actual"] - by_dept["budget"]) / by_dept["budget"] * 100).round(2)

# Category breakdown
by_cat = (
    df.groupby("category")
      .agg(actual=("actual","sum"))
      .reset_index()
      .sort_values("actual", ascending=False)
)

# Top KPIs
total_budget   = df["budget"].sum()
total_actual   = df["actual"].sum()
total_variance = total_actual - total_budget
overall_var_pct= total_variance / total_budget * 100
over_budget_depts = (by_dept["actual"] > by_dept["budget"]).sum()

print(f"   Total Budget : €{total_budget:,.0f}")
print(f"   Total Actual : €{total_actual:,.0f}")
print(f"   Variance     : €{total_variance:,.0f} ({overall_var_pct:.1f}%)")

# ── 3. CHARTS (saved as PNG for Excel embedding) ───────────────────────────────
print("📊 Creating charts …")

CHART_DIR = OUTPUT_DIR / "charts"
CHART_DIR.mkdir(exist_ok=True)

plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "axes.spines.top": False,
    "axes.spines.right": False,
    "axes.grid": True,
    "grid.alpha": 0.3,
    "grid.linestyle": "--",
})

# Chart 1: Monthly Budget vs Actual
fig, ax = plt.subplots(figsize=(10, 4))
x = np.arange(len(monthly))
w = 0.38
ax.bar(x - w/2, monthly["budget"]/1000,  w, label="Budget",  color=f"#{BLUE_MID}",  alpha=0.85)
ax.bar(x + w/2, monthly["actual"]/1000,  w, label="Actual",  color=f"#{ACCENT}",    alpha=0.85)
ax.set_xticks(x)
ax.set_xticklabels(monthly["month_str"], rotation=45, ha="right", fontsize=9)
ax.yaxis.set_major_formatter(mticker.FuncFormatter(lambda v,_: f"€{v:.0f}K"))
ax.set_title("Monthly Budget vs Actual (€K)", fontsize=13, fontweight="bold", pad=12)
ax.legend(framealpha=0)
plt.tight_layout()
chart1_path = CHART_DIR / "monthly_budget_actual.png"
fig.savefig(chart1_path, dpi=150, bbox_inches="tight")
plt.close()

# Chart 2: Variance % by Month
fig, ax = plt.subplots(figsize=(10, 3.5))
colors = [f"#{RED_SOFT}" if v > 0 else f"#{GREEN}" for v in monthly["var_pct"]]
ax.bar(monthly["month_str"], monthly["var_pct"], color=colors, alpha=0.85)
ax.axhline(0, color="black", linewidth=0.8)
ax.set_xticklabels(monthly["month_str"], rotation=45, ha="right", fontsize=9)
ax.yaxis.set_major_formatter(mticker.FuncFormatter(lambda v,_: f"{v:.0f}%"))
ax.set_title("Monthly Variance % (Over Budget = Red)", fontsize=13, fontweight="bold", pad=12)
plt.tight_layout()
chart2_path = CHART_DIR / "variance_pct.png"
fig.savefig(chart2_path, dpi=150, bbox_inches="tight")
plt.close()

# Chart 3: Department Actual Spending
fig, ax = plt.subplots(figsize=(8, 4))
bars = ax.barh(by_dept["department"], by_dept["actual"]/1000,
               color=[f"#{RED_SOFT}" if v > 0 else f"#{BLUE_MID}" for v in by_dept["var_pct"]],
               alpha=0.85)
ax.xaxis.set_major_formatter(mticker.FuncFormatter(lambda v,_: f"€{v:.0f}K"))
ax.set_title("Actual Spend by Department (€K)", fontsize=13, fontweight="bold", pad=12)
ax.invert_yaxis()
plt.tight_layout()
chart3_path = CHART_DIR / "dept_spending.png"
fig.savefig(chart3_path, dpi=150, bbox_inches="tight")
plt.close()

print("   ✅ Charts saved.")

# ── 4. BUILD EXCEL REPORT ─────────────────────────────────────────────────────
print("📝 Building Excel report …")

wb = openpyxl.Workbook()

# ── Helper styles ──────────────────────────────────────────────────────────────
def make_fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def thin_border():
    s = Side(style="thin", color=GREY_MID)
    return Border(left=s, right=s, top=s, bottom=s)

def header_font(color=WHITE, size=11, bold=True):
    return Font(name="Calibri", bold=bold, color=color, size=size)

def cell_font(bold=False, size=10, color="1E293B"):
    return Font(name="Calibri", bold=bold, size=size, color=color)

EUR = '#,##0.00 "€"'
PCT = '0.0"%"'

def set_col_widths(ws, widths: dict):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

# ══════════════════════════════════════════════════════════════════════════════
# SHEET 1 — EXECUTIVE SUMMARY
# ══════════════════════════════════════════════════════════════════════════════
ws1 = wb.active
ws1.title = "Executive Summary"
ws1.sheet_view.showGridLines = False

# Title banner
ws1.merge_cells("A1:H2")
ws1["A1"] = "FINANCIAL PERFORMANCE REPORT — 2024"
ws1["A1"].font      = Font(name="Calibri", bold=True, size=16, color=WHITE)
ws1["A1"].fill      = make_fill(BLUE_DARK)
ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")

ws1.merge_cells("A3:H3")
ws1["A3"] = f"Generated: {datetime.today().strftime('%d %B %Y')}  |  Fatih Demirtas — Automated Financial Reporting Pipeline"
ws1["A3"].font      = Font(name="Calibri", size=9, color="64748B", italic=True)
ws1["A3"].fill      = make_fill(GREY_LIGHT)
ws1["A3"].alignment = Alignment(horizontal="center", vertical="center")

ws1.row_dimensions[1].height = 30
ws1.row_dimensions[2].height = 30
ws1.row_dimensions[3].height = 16

# KPI boxes  (row 5)
kpi_data = [
    ("Total Budget",       f"€{total_budget:,.0f}",    BLUE_DARK),
    ("Total Actual",       f"€{total_actual:,.0f}",    BLUE_MID),
    ("Total Variance",     f"€{total_variance:+,.0f}", RED_SOFT if total_variance > 0 else GREEN),
    ("Variance %",         f"{overall_var_pct:+.1f}%", RED_SOFT if overall_var_pct > 0 else GREEN),
    ("Departments Over Budget", str(over_budget_depts), RED_SOFT if over_budget_depts > 0 else GREEN),
    ("Transactions",       f"{len(df):,}",             BLUE_DARK),
]

kpi_cols = [1, 2, 3, 4, 5, 6]   # A–F each span 1 col conceptually; we'll use A,B,C,D,E,F
for i, (label, value, color) in enumerate(kpi_data, start=1):
    col = get_column_letter(i)
    ws1.merge_cells(f"{col}5:{col}5")
    ws1.merge_cells(f"{col}6:{col}6")

    c_label = ws1[f"{col}5"]
    c_label.value     = label
    c_label.font      = Font(name="Calibri", size=9, bold=True, color=WHITE)
    c_label.fill      = make_fill(color)
    c_label.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    c_val = ws1[f"{col}6"]
    c_val.value     = value
    c_val.font      = Font(name="Calibri", size=12, bold=True, color=color)
    c_val.fill      = make_fill(GREY_LIGHT)
    c_val.alignment = Alignment(horizontal="center", vertical="center")

ws1.row_dimensions[5].height = 24
ws1.row_dimensions[6].height = 28

# Charts on summary sheet
img1 = XLImage(str(chart1_path))
img1.width, img1.height = 560, 220
ws1.add_image(img1, "A8")

img2 = XLImage(str(chart2_path))
img2.width, img2.height = 560, 200
ws1.add_image(img2, "A23")

img3 = XLImage(str(chart3_path))
img3.width, img3.height = 450, 230
ws1.add_image(img3, "A36")

set_col_widths(ws1, {l: 18 for l in ["A","B","C","D","E","F","G","H"]})

# ══════════════════════════════════════════════════════════════════════════════
# SHEET 2 — MONTHLY BREAKDOWN
# ══════════════════════════════════════════════════════════════════════════════
ws2 = wb.create_sheet("Monthly Breakdown")
ws2.sheet_view.showGridLines = False

headers = ["Month", "Budget (€)", "Actual (€)", "Variance (€)", "Variance %", "Transactions"]
header_fills = [BLUE_DARK, BLUE_MID, BLUE_MID, BLUE_MID, BLUE_MID, BLUE_MID]

for ci, (h, f) in enumerate(zip(headers, header_fills), start=1):
    c = ws2.cell(row=1, column=ci, value=h)
    c.font      = header_font()
    c.fill      = make_fill(f)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border    = thin_border()
ws2.row_dimensions[1].height = 22

for ri, row in monthly.iterrows():
    r = ri + 2
    vals = [row["month_str"], row["budget"], row["actual"], row["variance"], row["var_pct"]/100, row["txn"]]
    fmts = [None, EUR, EUR, EUR, '0.0%', None]
    for ci, (v, fmt) in enumerate(zip(vals, fmts), start=1):
        c = ws2.cell(row=r, column=ci, value=v)
        c.font      = cell_font()
        c.alignment = Alignment(horizontal="center")
        c.border    = thin_border()
        if fmt:
            c.number_format = fmt
        # colour variance cells
        if ci == 4:
            c.font = cell_font(bold=True, color=RED_SOFT if v > 0 else GREEN)
        if ri % 2 == 0:
            c.fill = make_fill(GREY_LIGHT)

set_col_widths(ws2, {"A":14,"B":16,"C":16,"D":16,"E":14,"F":14})

# ══════════════════════════════════════════════════════════════════════════════
# SHEET 3 — DEPARTMENT ANALYSIS
# ══════════════════════════════════════════════════════════════════════════════
ws3 = wb.create_sheet("Department Analysis")
ws3.sheet_view.showGridLines = False

headers3 = ["Department", "Budget (€)", "Actual (€)", "Variance (€)", "Variance %", "Status", "Transactions"]
for ci, h in enumerate(headers3, start=1):
    c = ws3.cell(row=1, column=ci, value=h)
    c.font      = header_font()
    c.fill      = make_fill(BLUE_DARK)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border    = thin_border()
ws3.row_dimensions[1].height = 22

for ri, row in by_dept.iterrows():
    r = ri + 2
    over = row["actual"] > row["budget"]
    status = "⚠ Over Budget" if over else "✅ On Track"
    vals = [row["department"], row["budget"], row["actual"],
            row["actual"]-row["budget"], row["var_pct"]/100, status, row["txn"]]
    fmts = [None, EUR, EUR, EUR, '0.0%', None, None]
    for ci, (v, fmt) in enumerate(zip(vals, fmts), start=1):
        c = ws3.cell(row=r, column=ci, value=v)
        c.font      = cell_font()
        c.alignment = Alignment(horizontal="center")
        c.border    = thin_border()
        if fmt:
            c.number_format = fmt
        if ci == 6:
            c.font = cell_font(bold=True, color=RED_SOFT if over else GREEN)
        if ri % 2 == 0:
            c.fill = make_fill(GREY_LIGHT)

# Department chart inside sheet
img3b = XLImage(str(chart3_path))
img3b.width, img3b.height = 480, 260
ws3.add_image(img3b, "A12")

set_col_widths(ws3, {"A":18,"B":16,"C":16,"D":16,"E":14,"F":16,"G":14})

# ══════════════════════════════════════════════════════════════════════════════
# SHEET 4 — RAW DATA
# ══════════════════════════════════════════════════════════════════════════════
ws4 = wb.create_sheet("Raw Data")
ws4.sheet_view.showGridLines = False

raw_cols = ["date","department","category","budget","actual","variance","var_pct","over_budget"]
display_headers = ["Date","Department","Category","Budget (€)","Actual (€)","Variance (€)","Var %","Over Budget"]

for ci, h in enumerate(display_headers, start=1):
    c = ws4.cell(row=1, column=ci, value=h)
    c.font      = header_font()
    c.fill      = make_fill(BLUE_MID)
    c.alignment = Alignment(horizontal="center")
    c.border    = thin_border()
ws4.row_dimensions[1].height = 20

for ri, row in df[raw_cols].iterrows():
    r = ri + 2
    vals = [row["date"].strftime("%Y-%m-%d"), row["department"], row["category"],
            row["budget"], row["actual"], row["variance"], row["var_pct"]/100,
            "Yes" if row["over_budget"] else "No"]
    fmts = [None, None, None, EUR, EUR, EUR, '0.0%', None]
    for ci, (v, fmt) in enumerate(zip(vals, fmts), start=1):
        c = ws4.cell(row=r, column=ci, value=v)
        c.font      = cell_font(size=9)
        c.alignment = Alignment(horizontal="center")
        c.border    = thin_border()
        if fmt:
            c.number_format = fmt
        if ri % 2 == 0:
            c.fill = make_fill(GREY_LIGHT)

ws4.auto_filter.ref = f"A1:{get_column_letter(len(display_headers))}1"
set_col_widths(ws4, {"A":13,"B":14,"C":13,"D":14,"E":14,"F":14,"G":10,"H":12})

# ── Save ──────────────────────────────────────────────────────────────────────
wb.save(REPORT_FILE)
print(f"\n✅ Report saved → {REPORT_FILE}")
print("   Sheets: Executive Summary | Monthly Breakdown | Department Analysis | Raw Data")
